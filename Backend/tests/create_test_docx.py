from docx import Document


document = Document()

document.add_heading(
    "Enterprise AI Chatbot DOCX Test",
    level=1,
)

document.add_paragraph(
    "This is a test document for the DOCX extraction pipeline."
)

document.add_paragraph(
    "The DOCX parser should extract this content correctly."
)

document.add_paragraph(
    "This document will later be processed by the text cleaner."
)

document.save(
    "storage/test_documents/test_document.docx"
)

print("Test DOCX created successfully.")
